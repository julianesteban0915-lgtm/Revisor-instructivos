import io
import pdfplumber
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
import requests

# ----------------------------------------------------------------------
# Revisor de Instructivos de Embalaje — Garces Fruit
# App web: el usuario sube el PDF del instructivo y obtiene la revision
# contra el kit base maestro. Se ignoran las columnas Observacion y Pallet.
# ----------------------------------------------------------------------

st.set_page_config(page_title="Revisor de Instructivos", page_icon="🍒", layout="centered")

# ---------------- helpers ----------------
def norm(s):
    return "" if s is None else str(s).strip().upper()

def code(s):
    n = norm(s)
    return n.split()[0] if n else ""

def numv(s):
    d = "".join(ch for ch in str("" if s is None else s) if ch.isdigit())
    return int(d) if d else None

# ---------------- kit base desde GitHub ----------------
GITHUB_EXCEL_URL = "https://raw.githubusercontent.com/julianesteban0915-lgtm/Revisor-instructivos/main/KIT_EMBALAJE_TODOS.xlsx"

@st.cache_data(show_spinner="Cargando kit base desde GitHub...", ttl=300)
def load_base_github():
    """Carga el Excel de kit base directamente desde GitHub."""
    try:
        r = requests.get(GITHUB_EXCEL_URL, timeout=15)
        if r.status_code == 200:
            return load_base(r.content), None
        else:
            return [], f"No se pudo descargar el kit base (código {r.status_code})"
    except Exception as e:
        return [], str(e)

@st.cache_data(show_spinner=False)
def load_base(file_bytes):
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    base = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        hi = env = emb = eti = caj = -1
        for i, r in enumerate(rows[:15]):
            u = [norm(c).replace(" ", "") for c in r]

            def find(keys):
                for j, c in enumerate(u):
                    if any(k in c for k in keys):
                        return j
                return -1

            e, b, t = find(["ENVASE"]), find(["EMBALAJE"]), find(["ETIQUETA"])
            if e >= 0 and b >= 0 and t >= 0:
                hi, env, emb, eti = i, e, b, t
                caj = find(["CANTIDADDECAJAS", "CAJASPALLET", "CAJAS", "CANTIDAD"])
                break
        if hi < 0:
            continue
        for r in rows[hi + 1:]:
            ev = code(r[env]) if env < len(r) else ""
            if not ev:
                continue
            base.append({
                "envase": ev,
                "embalaje": code(r[emb]) if emb < len(r) else "",
                "etiqueta": norm(r[eti]) if eti < len(r) else "",
                "cajas": numv(r[caj]) if 0 <= caj < len(r) else None,
                "hoja": sn,
            })
    return base

# ---------------- instructivo PDF ----------------
def extract_kits(pdf_bytes):
    kits = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            for t in page.extract_tables():
                if len(t) < 2:
                    continue
                hdr = [norm(c).replace(" ", "") for c in t[1]]
                if "ENVASE" in hdr and "EMBALAJE" in hdr and "ETIQUETA" in hdr:
                    iE, iB, iT = hdr.index("ENVASE"), hdr.index("EMBALAJE"), hdr.index("ETIQUETA")
                    iCaj = hdr.index("CAJAS/PALLET") if "CAJAS/PALLET" in hdr else -1
                    iCal = hdr.index("CALIBRES") if "CALIBRES" in hdr else (
                        hdr.index("CALIBRE") if "CALIBRE" in hdr else -1)
                    for r in t[2:]:
                        ev = code(r[iE]) if iE < len(r) else ""
                        if not ev:
                            continue
                        kits.append({
                            "envase": ev,
                            "embalaje": code(r[iB]) if iB < len(r) else "",
                            "etiqueta": norm(r[iT]) if iT < len(r) else "",
                            "calibre": norm(r[iCal]) if 0 <= iCal < len(r) else "",
                            "cajas": numv(r[iCaj]) if 0 <= iCaj < len(r) else None,
                        })
    return kits

# ---------------- comparacion ----------------
def review(base, kits):
    out = []
    for k in kits:
        m = [b for b in base if b["envase"] == k["envase"]
             and b["embalaje"] == k["embalaje"] and b["etiqueta"] == k["etiqueta"]]
        if not m:
            partial = [b for b in base if b["envase"] == k["envase"] and b["embalaje"] == k["embalaje"]]
            if partial:
                sug = ", ".join(sorted(set(b["etiqueta"] for b in partial))[:3])
                st_, note = "ERROR", f'Etiqueta "{k["etiqueta"] or "(vacía)"}" no existe para este envase/embalaje. La base registra: {sug}.'
            else:
                st_, note = "ERROR", "La combinación Envase + Embalaje no existe en el kit base."
        else:
            cb = sorted(set(b["cajas"] for b in m if b["cajas"] is not None))
            if k["cajas"] is None or not cb or k["cajas"] in cb:
                st_, note = "OK", ""
            else:
                st_, note = "ADVERTENCIA", f'Cajas/pallet = {k["cajas"]}; la base registra {" / ".join(map(str, cb))} para este kit.'
        out.append({**k, "estado": st_, "detalle": note})
    return out

# ---------------- UI ----------------
st.markdown("""
<style>
#MainMenu, footer, header, [data-testid="stToolbar"], [data-testid="stDecoration"], [data-testid="stStatusWidget"]{visibility:hidden; height:0;}
.block-container{max-width:840px; padding-top:1.4rem;}
.gf-header{background:#2f5d2a; border-radius:14px; padding:18px 22px; margin-bottom:6px; display:flex; align-items:center; gap:14px;}
.gf-logo{width:44px; height:44px; border-radius:11px; background:#ffffff; color:#2f5d2a; display:flex; align-items:center; justify-content:center; font-weight:700; font-size:19px; flex-shrink:0; font-family:sans-serif;}
.gf-header h1{margin:0; color:#ffffff; font-size:22px; font-weight:600; line-height:1.15;}
.gf-header p{margin:2px 0 0; color:#cfe0c6; font-size:12.5px;}
.res{border:1px solid #e3e2d8;border-radius:12px;padding:13px 16px;margin-bottom:9px}
.res.OK{border-left:5px solid #97c459}
.res.ERROR{border-left:5px solid #f09595;background:#fefafa}
.res.ADVERTENCIA{border-left:5px solid #ef9f27;background:#fffdf8}
.rtop{display:flex;justify-content:space-between;align-items:center;margin-bottom:6px}
.rid{font-family:monospace;font-size:13px;font-weight:600}
.tag{font-size:11px;font-weight:700;padding:3px 11px;border-radius:20px}
.tag.OK{background:#eaf3de;color:#3b6d11}
.tag.ERROR{background:#fcebeb;color:#a32d2d}
.tag.ADVERTENCIA{background:#faeeda;color:#854f0b}
.rfields{font-family:monospace;font-size:12.5px;color:#444}
.note{font-size:12.5px;margin-top:7px;color:#a32d2d}
.note.ADVERTENCIA{color:#854f0b}
.kit-info{background:#f0f7ee;border-radius:8px;padding:8px 12px;font-size:12px;color:#2f5d2a;margin-bottom:4px;}
</style>
""", unsafe_allow_html=True)

st.markdown(
    '<div class="gf-header"><div class="gf-logo">GF</div>'
    '<div><h1>Revisor de Instructivos</h1>'
    '<p>Garces Fruit · revisión de kits contra el kit base maestro</p></div></div>',
    unsafe_allow_html=True)
st.caption("Sube el instructivo en PDF y obtén la revisión. "
           "Se ignoran las columnas **Observación** y **Pallet**.")

# --- kit base desde GitHub (automático) ---
with st.expander("⚙️ Kit base maestro", expanded=False):
    base_github, error_github = load_base_github()

    if error_github:
        st.warning(f"⚠️ No se pudo cargar desde GitHub: {error_github}")
        st.info("Sube el Excel manualmente como alternativa.")
        base = []
    else:
        hojas = sorted(set(b["hoja"] for b in base_github))
        st.success(f"✅ Kit base cargado desde GitHub: **{len(base_github)} kits** en **{len(hojas)} hoja(s)**.")
        st.markdown(f'<div class="kit-info">📋 Hojas: {" · ".join(hojas)}</div>', unsafe_allow_html=True)
        st.caption("Se actualiza automáticamente cuando subes un Excel nuevo a GitHub.")
        base = base_github

    st.divider()
    st.markdown("**¿Actualizaste el Excel?** Súbelo a GitHub:")
    st.code("github.com/julianesteban0915-lgtm/Revisor-instructivos → Subir archivo → KIT_EMBALAJE_TODOS.xlsx", language=None)

    st.markdown("**O reemplaza manualmente solo para esta sesión:**")
    up_base = st.file_uploader("Subir Excel kit base (solo esta sesión)", type=["xlsx", "xlsm", "xls"], key="base")
    if up_base is not None:
        base = load_base(up_base.read())
        hojas_manual = sorted(set(b["hoja"] for b in base))
        st.success(f"Kit base manual cargado: {len(base)} kits en {len(hojas_manual)} hoja(s).")

if base:
    st.caption(f"Comparando contra **{len(base)} kits** del kit base.")

# --- subir instructivo ---
st.subheader("Sube el instructivo (PDF)")
up_pdf = st.file_uploader("Instructivo de embalaje", type=["pdf"], key="pdf",
                          label_visibility="collapsed")

if up_pdf is not None and base:
    kits = extract_kits(up_pdf.read())
    if not kits:
        st.error("No pude leer la tabla de órdenes específicas en este PDF. "
                 "¿Es un instructivo en formato GF-IND-PL-003?")
    else:
        res = review(base, kits)
        n_ok  = sum(r["estado"] == "OK" for r in res)
        n_er  = sum(r["estado"] == "ERROR" for r in res)
        n_wa  = sum(r["estado"] == "ADVERTENCIA" for r in res)

        c1, c2, c3 = st.columns(3)
        c1.metric("OK", n_ok)
        c2.metric("Errores", n_er)
        c3.metric("Advertencias", n_wa)

        st.divider()
        for r in res:
            fields = f'{r["envase"]} · {r["embalaje"]} · {r["etiqueta"] or "—"} · cajas/pallet: {r["cajas"] if r["cajas"] is not None else "—"}'
            note = f'<div class="note {r["estado"]}">{r["detalle"]}</div>' if r["detalle"] else ""
            st.markdown(
                f'<div class="res {r["estado"]}"><div class="rtop">'
                f'<span class="rid">{r["envase"]} · {r["embalaje"]} · {r["etiqueta"] or "—"}</span>'
                f'<span class="tag {r["estado"]}">{r["estado"]}</span></div>'
                f'<div class="rfields">{fields}</div>{note}</div>',
                unsafe_allow_html=True)

        # exportar a Excel
        df = pd.DataFrame([{
            "Envase": r["envase"], "Embalaje": r["embalaje"], "Etiqueta": r["etiqueta"],
            "Calibre": r["calibre"], "Cajas/pallet": r["cajas"],
            "Estado": r["estado"], "Detalle": r["detalle"],
        } for r in res])
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as w:
            df.to_excel(w, index=False, sheet_name="Revisión")
        st.download_button("⬇ Descargar resultado en Excel", buf.getvalue(),
                           file_name="revision_instructivo.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

elif up_pdf is not None and not base:
    st.warning("No se pudo cargar el kit base. Sube el Excel manualmente arriba.")
